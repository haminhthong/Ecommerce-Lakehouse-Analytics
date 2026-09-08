"""Tạo hướng dẫn cải thiện dự án theo bốn tầng đánh giá."""

from pathlib import Path

from build_project_improvement_docx import (
    BLUE,
    DARK_BLUE,
    GOLD,
    LIGHT_BLUE,
    NAVY,
    RED,
    add_bullet,
    add_callout,
    add_contents,
    add_number,
    add_table,
    configure_document,
    set_run_font,
)
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Huong_dan_cai_tien_du_an_theo_4_tang.docx"


def cover(document):
    document.add_paragraph().paragraph_format.space_after = Pt(44)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(kicker.add_run("GLOBALCART INTELLIGENCE"), 11, True, BLUE)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run("HƯỚNG DẪN CẢI TIẾN DỰ ÁN\nTHEO BỐN TẦNG"), 28, True, NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        subtitle.add_run(
            "Problem → AI/ML Correctness → Software Engineering\n→ Production & Business Value"
        ),
        14,
        False,
        DARK_BLUE,
    )
    document.add_paragraph().paragraph_format.space_after = Pt(28)
    add_callout(
        document,
        "Mục tiêu",
        "Chuyển repository hiện tại thành một dự án Data Engineering/BI trung thực, có thể kiểm chứng, dễ chạy và đủ sức thuyết phục khi đưa vào CV.",
        fill="EAF2F8",
    )
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(36)
    set_run_font(meta.add_run("Bản hướng dẫn chi tiết • Cập nhật 01/09/2026"), 10, False, "667085")
    document.add_page_break()


def checklist_table(document):
    document.add_heading("Kết quả đánh giá nhanh", level=1)
    add_table(
        document,
        ["Tầng", "Trạng thái", "Kết luận"],
        [
            (
                "Problem",
                "Đạt",
                "Bài toán Data Engineering/BI hợp lý; cần giảm tuyên bố Enterprise.",
            ),
            (
                "AI/ML correctness",
                "Không áp dụng cho ML",
                "RFM/ABC là rule-based analytics, không có training/inference.",
            ),
            (
                "Software Engineering",
                "Đạt một phần",
                "Ruff và 12 test đạt; thiếu edge/integration tests.",
            ),
            (
                "Production/Business",
                "Chưa đạt",
                "Chưa có load test, auth/TLS, monitoring và end-to-end evidence.",
            ),
        ],
        [2100, 2200, 5060],
    )
    add_callout(
        document,
        "Nguyên tắc",
        "Không thêm ML chỉ để làm dự án trông phức tạp. Hoàn thiện correctness, integration và demo trước; ML chỉ được thêm khi có mục tiêu dự đoán, nhãn và cách đánh giá rõ ràng.",
        fill="FFF8E8",
        color=GOLD,
    )


def problem_layer(document):
    document.add_heading("1. Tầng Problem — Giải quyết đúng bài toán", level=1)
    document.add_paragraph(
        "Bài toán đúng của dự án là biến dữ liệu giao dịch ecommerce thô thành Lakehouse có kiểm soát chất lượng, mô hình sao và data marts phục vụ Power BI/MongoDB. Đây là bài toán Data Engineering và BI Engineering, không phải AI product."
    )
    document.add_heading("1.1 Chốt problem statement", level=2)
    add_callout(
        document,
        "Câu mô tả khuyến nghị",
        "GlobalCart Intelligence là portfolio project mô phỏng nền tảng batch Ecommerce Lakehouse bằng Spark và Delta Lake, cung cấp dữ liệu đã kiểm soát cho Power BI và MongoDB.",
        fill=LIGHT_BLUE,
    )
    document.add_heading("1.2 Xác định grain và KPI", level=2)
    add_bullet(document, "Một dòng Silver/FactSales = một sản phẩm trong một đơn hàng.")
    add_bullet(document, "Order_ID có thể lặp; KPI số đơn phải dùng countDistinct(Order_ID).")
    add_bullet(document, "Số dòng giao dịch và số đơn hàng là hai KPI khác nhau.")
    add_bullet(
        document, "Profit Margin phải tính SUM(Profit)/SUM(Revenue), không trung bình tỷ lệ dòng."
    )
    document.add_heading("Thay đổi cần thực hiện", level=3)
    add_number(document, "Đổi các mart đang dùng count(Order_ID) sang countDistinct(Order_ID).")
    add_number(document, "Thêm test một đơn có nhiều sản phẩm: 3 dòng nhưng chỉ 2 đơn.")
    add_number(document, "Ghi grain ngay cạnh schema FactSales và trong data dictionary.")
    document.add_heading("1.3 Hoàn thiện data contract", level=2)
    add_table(
        document,
        ["Rule", "Cách kiểm tra", "Xử lý khi lỗi"],
        [
            ("Required fields", "Order/Customer/Product/Date không null", "Đưa vào quarantine"),
            ("Numeric range", "Quantity > 0, Price ≥ 0, Discount ∈ [0,1]", "Từ chối bản ghi"),
            ("Date consistency", "Year/Month khớp Order_Date", "Ghi reason code"),
            ("Revenue formula", "|actual - expected| ≤ tolerance", "Cảnh báo hoặc quarantine"),
            ("Allowed values", "Status/Payment/Shipping thuộc domain", "Từ chối hoặc map Unknown"),
        ],
        [2200, 4300, 2860],
    )
    document.add_heading("Acceptance criteria", level=3)
    add_bullet(document, "Mỗi rule có mã, mô tả, severity và test biên.")
    add_bullet(document, "Invalid records không bị xóa âm thầm.")
    add_bullet(document, "Báo cáo số dòng Raw, Valid, Invalid và tỷ lệ reject.")


def analytics_layer(document):
    document.add_heading("2. Tầng AI/ML Correctness", level=1)
    add_callout(
        document,
        "Kết luận",
        "Dự án hiện không có mô hình ML. Data leakage, train/validation/test và training/inference preprocessing được đánh dấu Không áp dụng. RFM và ABC là phân tích theo quy tắc.",
        fill="EAF2F8",
    )
    document.add_heading("2.1 Đồng nhất RFM", level=2)
    document.add_paragraph(
        "Hiện Pandas dùng qcut theo phân vị, còn Spark dùng ngưỡng ngày/tần suất cố định. Một khách hàng có thể nhận segment khác nhau giữa báo cáo portfolio và Gold Mart."
    )
    add_table(
        document,
        ["Lựa chọn", "Ưu điểm", "Điều kiện"],
        [
            (
                "Quantile scoring",
                "Tự thích nghi phân phối dữ liệu",
                "Pandas và Spark phải dùng cùng cách xử lý tie/bin",
            ),
            (
                "Business thresholds",
                "Dễ giải thích và ổn định",
                "Ngưỡng phải nằm trong config và có chủ sở hữu nghiệp vụ",
            ),
        ],
        [2400, 3300, 3660],
    )
    document.add_heading("Cách triển khai khuyến nghị", level=3)
    add_number(document, "Tạo analytics_rules.py chứa tên segment, cutoff và version của rule.")
    add_number(document, "Chọn một thuật toán duy nhất; không duy trì hai business rule khác nhau.")
    add_number(document, "Truyền analysis_date/cutoff thay vì luôn lấy max date ngầm định.")
    add_number(document, "Viết contract test chạy cùng fixture qua Pandas và Spark rồi so segment.")
    document.add_heading("2.2 Chuẩn hóa ABC", level=2)
    document.add_paragraph(
        "Nên phân loại bằng cumulative revenue trước sản phẩm hiện tại. Sản phẩm làm tổng vượt 80% vẫn thuộc A; sản phẩm làm tổng vượt 95% vẫn thuộc B. Cùng rule phải được dùng trong Pandas và Spark."
    )
    add_table(
        document,
        ["Edge case", "Kết quả mong đợi"],
        [
            ("Sản phẩm đầu tiên chiếm >80%", "Vẫn thuộc A và rule được giải thích"),
            ("Đúng ngưỡng 80%", "Kết quả nhất quán theo quy ước đã chọn"),
            ("Tổng revenue = 0", "Không chia cho 0; trả bảng hợp lệ hoặc lỗi rõ ràng"),
            ("Revenue âm", "Bị data quality gate chặn trước ABC"),
            ("Dataset rỗng", "Trả schema rỗng hoặc ValueError có thông báo"),
        ],
        [3500, 5860],
    )
    document.add_heading("2.3 Khi nào mới thêm ML", level=2)
    add_bullet(document, "Return/cancellation prediction: có target rõ và split theo thời gian.")
    add_bullet(document, "Revenue forecasting: có naive seasonal baseline và backtesting.")
    add_bullet(
        document,
        "Churn prediction: định nghĩa churn/cutoff/observation window trước khi tạo feature.",
    )
    add_bullet(document, "Chỉ dùng test set một lần để báo cáo cuối; preprocessing fit trên train.")


def engineering_layer(document):
    document.add_heading("3. Tầng Software Engineering", level=1)
    document.add_heading("3.1 Mở rộng test edge cases", level=2)
    add_table(
        document,
        ["Nhóm", "Test cần bổ sung"],
        [
            (
                "Data quality",
                "Empty CSV, null required fields, invalid date, negative metrics, Unicode, duplicates",
            ),
            (
                "RFM",
                "Một customer, identical values, insufficient qcut bins, future date, explicit cutoff",
            ),
            ("ABC", "Zero total, exact 80/95%, crossing products, ties, empty input"),
            (
                "MongoDB",
                "Connection failure, partial batch failure, BSON conversion, missing index field",
            ),
            ("CLI", "Missing file, child exit != 0, invalid command, report failure"),
            ("Spark", "Schema, grain, foreign-key null, idempotency, local Delta read/write"),
        ],
        [2100, 7260],
    )
    document.add_heading("3.2 Tách Spark main()", level=2)
    add_bullet(document, "ingestion.py: đọc nguồn và schema.")
    add_bullet(document, "silver.py: clean, cast, enrich và quality gate.")
    add_bullet(document, "dimensions.py: DimDate và các dimensions.")
    add_bullet(document, "marts.py: KPI/RFM/ABC Spark transformations.")
    add_bullet(document, "storage.py: save/register/check Delta.")
    add_bullet(document, "pipeline.py: chỉ orchestration và lifecycle Spark.")
    document.add_heading("Quy tắc refactor", level=3)
    add_number(document, "Tách từng nhóm nhỏ, không di chuyển 800 dòng trong một commit.")
    add_number(document, "Thêm regression test trước khi tách transformation.")
    add_number(document, "Giữ schema public và table names để không phá Power BI.")
    add_number(document, "Chạy Ruff, unit tests và integration fixture sau mỗi bước.")
    document.add_heading("3.3 Portability", level=2)
    add_bullet(
        document,
        "Thay path C:\\Users\\... trong EcommerceDW.sql bằng SQLCMD variable hoặc import script.",
    )
    add_bullet(document, "Tạo requirements lock để CI và máy người dùng dùng cùng versions.")
    add_bullet(document, "Thực hiện clean-room test trên VM/máy mới.")
    add_bullet(document, "PBIX cần hướng dẫn đổi data source và ODBC DSN.")
    document.add_heading("3.4 Logging và error handling", level=2)
    add_bullet(document, "Thay print vận hành bằng LOGGER; print chỉ dành cho CLI output.")
    add_bullet(document, "Log run_id, stage, row_count, invalid_count, duration và Delta version.")
    add_bullet(document, "Bắt exception cụ thể; nếu chỉ log thì raise lại với nguyên nhân gốc.")


def production_layer(document):
    document.add_heading("4. Tầng Production & Business Value", level=1)
    document.add_heading("4.1 Chứng minh end-to-end", level=2)
    for item in [
        "Chạy doctor và lưu dependency report.",
        "Chạy check: data contract, report generator và test suite.",
        "Chạy local pipeline hai lần để chứng minh idempotency.",
        "Kiểm tra row counts, schema và foreign keys từng layer.",
        "Đồng bộ MongoDB; đối chiếu document counts và indexes.",
        "Khởi động Thrift; refresh Power BI và chụp evidence.",
    ]:
        add_number(document, item)
    document.add_heading("4.2 Incremental và performance", level=2)
    add_bullet(document, "Xác định Order_Line_ID/business key trước khi Delta MERGE.")
    add_bullet(document, "Partition FactSales theo Year/Month chỉ khi dữ liệu đủ lớn.")
    add_bullet(document, "Gom Spark actions để tránh nhiều count/collect.")
    add_bullet(
        document,
        "Benchmark cùng dữ liệu, Spark config và phần cứng; ghi median của nhiều lần chạy.",
    )
    document.add_heading("4.3 100 users", level=2)
    add_callout(
        document,
        "Không nên tuyên bố",
        "Spark local + Derby + Thrift local chưa có bằng chứng phục vụ 100 concurrent users. Với portfolio, nên dùng Power BI Import Mode và để Power BI Service phục vụ người xem.",
        fill="FFF0F0",
        color=RED,
    )
    add_bullet(
        document,
        "Nếu DirectQuery: cần cluster/SQL warehouse, external metastore, auth, cache và load test.",
    )
    add_bullet(
        document, "Định nghĩa SLA: refresh time, query latency, concurrency và failure rate."
    )
    document.add_heading("4.4 Security và privacy", level=2)
    add_bullet(document, "MongoDB authentication/TLS và least-privilege accounts.")
    add_bullet(document, "Thrift authentication, TLS và network allowlist.")
    add_bullet(document, "Secret manager; không log connection string có credentials.")
    add_bullet(document, "Ghi nguồn/license dataset, synthetic/real và chính sách PII.")
    add_bullet(
        document, "Mask Customer_ID khi publish public nếu dữ liệu không hoàn toàn synthetic."
    )


def portfolio_layer(document):
    document.add_heading("5. README, Demo và CV", level=1)
    document.add_heading("5.1 Sửa tuyên bố", level=2)
    add_table(
        document,
        ["Không nên viết", "Nên viết"],
        [
            ("Enterprise production-ready", "Portfolio project mô phỏng kiến trúc thực tế"),
            ("100% Pytest coverage", "12 automated tests cho core logic"),
            ("Tối ưu tốc độ", "Thiết kế data marts để giảm tính toán dashboard; chưa benchmark"),
            ("AI-powered RFM/ABC", "Rule-based RFM and Pareto ABC analytics"),
        ],
        [4300, 5060],
    )
    document.add_heading("5.2 Demo cần có", level=2)
    add_bullet(document, "Chỉ giữ powerbi/GlobalCart_Analytics.pbix.")
    add_bullet(document, "Thêm ảnh Executive Overview, RFM, ABC và Fulfillment/Returns.")
    add_bullet(document, "Nêu ba insight, một khuyến nghị và một limitation.")
    add_bullet(document, "Có log hoặc video chứng minh refresh từ Gold thành công.")
    document.add_heading("5.3 Limitations bắt buộc", level=2)
    add_bullet(document, "Dataset nhỏ và batch; chưa có streaming/CDC.")
    add_bullet(document, "RFM/ABC là rule-based, không phải ML.")
    add_bullet(document, "Chưa benchmark cluster hoặc load test concurrency.")
    add_bullet(document, "Derby và local Thrift chỉ dành cho demo.")
    add_bullet(document, "Integration tests yêu cầu Spark/Delta/HDFS/MongoDB.")


def roadmap(document):
    document.add_heading("6. Roadmap ưu tiên", level=1)
    add_table(
        document,
        ["Ưu tiên", "Công việc", "Tiêu chí hoàn thành"],
        [
            (
                "P0",
                "countDistinct, RFM/ABC thống nhất, SQL path",
                "Contract tests Pandas/Spark đạt",
            ),
            ("P1", "Edge tests và Spark integration", "Local E2E chạy hai lần không sai lệch"),
            (
                "P2",
                "README claims, limitations, PBIX screenshots",
                "Recruiter hiểu dự án trong 2 phút",
            ),
            (
                "P3",
                "Incremental MERGE, partition, external metastore",
                "Benchmark và runbook có số liệu",
            ),
            ("P4", "Auth/TLS, monitoring, load test", "Đạt SLA được định nghĩa trước"),
        ],
        [1000, 4300, 4060],
    )
    document.add_heading("Definition of Done", level=2)
    checks = [
        "Ruff và toàn bộ unit/edge tests pass.",
        "Pandas/Spark RFM và ABC cho cùng kết quả trên contract fixture.",
        "Không còn path máy cá nhân hoặc secret commit.",
        "Local và HDFS pipeline có evidence chạy thành công.",
        "MongoDB sync và Power BI refresh được đối chiếu số dòng.",
        "README chỉ chứa claim có thể kiểm chứng.",
        "Có Dataset Source/License, Security và Limitations.",
    ]
    for check in checks:
        add_bullet(document, f"[ ] {check}")


def build():
    document = Document()
    configure_document(document)
    cover(document)
    add_contents(document)
    checklist_table(document)
    problem_layer(document)
    analytics_layer(document)
    engineering_layer(document)
    production_layer(document)
    portfolio_layer(document)
    roadmap(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
