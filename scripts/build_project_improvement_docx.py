"""Tạo báo cáo Word giải thích quá trình cải thiện GlobalCart Intelligence."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Bao_cao_cai_tien_du_an_GlobalCart_Intelligence.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
RED = "9B1C1C"
GOLD = "7A5A00"


def set_run_font(run, size=11, bold=False, color="000000", italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for item_width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(item_width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.5, bold=True, color=NAVY)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_run_font(run, size=9.3)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def add_callout(doc, label, text, fill=LIGHT_GRAY, color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, bold=True, color=color)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, color=color)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "GLOBALCART INTELLIGENCE  |  PROJECT IMPROVEMENT REPORT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=8.5, bold=True, color=MID_GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("Trang ")
    set_run_font(footer_run, size=9, color=MID_GRAY)
    add_page_number(footer)


def add_cover(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(50)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("BÁO CÁO KỸ THUẬT & PORTFOLIO")
    set_run_font(run, size=11, bold=True, color=BLUE)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("PHÂN TÍCH, CẢI TIẾN VÀ\nHOÀN THIỆN DỰ ÁN")
    set_run_font(run, size=27, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("GlobalCart Intelligence\nEnd-to-End Ecommerce Lakehouse & BI Platform")
    set_run_font(run, size=15, color=DARK_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(32)
    add_callout(
        doc,
        "Mục tiêu",
        "Giải thích toàn bộ kiến trúc, luồng dữ liệu, vấn đề kỹ thuật, các thay đổi đã thực hiện và lộ trình đưa dự án lên GitHub/CV.",
        fill="EAF2F8",
    )
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(36)
    run = meta.add_run("Vai trò định hướng: Data Engineer / BI Engineer\nNgày cập nhật: 30/08/2026")
    set_run_font(run, size=10.5, color=MID_GRAY)
    doc.add_page_break()


def add_contents(doc):
    doc.add_heading("Mục lục nội dung", level=1)
    items = [
        "1. Tóm tắt dự án và định hướng cải thiện",
        "2. Kiến trúc và luồng hoạt động",
        "3. Đánh giá kiến trúc hiện tại",
        "4. Technical debt và code smell",
        "5. Kế hoạch cải tiến theo giai đoạn",
        "6. Những thay đổi đã thực hiện",
        "7. Data quality, testing và CI",
        "8. Tối ưu hiệu năng và bảo mật",
        "9. Cách cài đặt và chạy dự án",
        "10. Kết quả, hạn chế và bước tiếp theo",
        "11. Cách trình bày dự án trong CV/phỏng vấn",
        "12. Checklist bàn giao",
    ]
    for item in items:
        add_bullet(doc, item)
    add_callout(
        doc,
        "Phạm vi",
        "Báo cáo dựa trên trạng thái repository hiện tại. Integration test Spark–HDFS–Delta–MongoDB cần các dịch vụ tương ứng hoạt động trên máy chạy.",
        fill="FFF8E8",
        color=GOLD,
    )
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)

    doc.add_heading("1. Tóm tắt dự án và định hướng cải thiện", level=1)
    doc.add_paragraph(
        "GlobalCart Intelligence là dự án Data Engineering và Business Intelligence mô phỏng nền tảng phân tích thương mại điện tử toàn cầu. Hệ thống nhận dữ liệu giao dịch CSV, lưu trữ trên HDFS, xử lý bằng Spark/Delta Lake theo kiến trúc Medallion, tạo mô hình sao và data marts, sau đó phục vụ Power BI và MongoDB."
    )
    add_table(
        doc,
        ["Thành phần", "Hiện trạng"],
        [
            ("Input", "CSV gồm 10.000 dòng giao dịch và 26 thuộc tính"),
            ("Processing", "PySpark, HDFS và Delta Lake"),
            ("Modeling", "1 fact, 7 dimensions và 10 data marts"),
            ("Serving", "Hive/Thrift cho Power BI; Delta sang MongoDB theo batch"),
            ("Quality", "Data contract, quality gate, pytest, Ruff và GitHub Actions"),
            ("Portfolio", "README, kiến trúc, data dictionary, business insights và CV bullets"),
        ],
        [2200, 7160],
    )
    doc.add_heading("Định hướng đề xuất", level=2)
    add_bullet(
        doc,
        "Giữ tên thương hiệu: GlobalCart Intelligence — End-to-End Ecommerce Lakehouse & BI Platform.",
    )
    add_bullet(
        doc, "Đặt Delta Lake làm single source of truth để Power BI và MongoDB không lệch số liệu."
    )
    add_bullet(
        doc,
        "Tập trung câu chuyện CV vào ingestion, data quality, dimensional modeling, serving và CI.",
    )
    add_bullet(doc, "Không tuyên bố hiệu năng production khi chưa có benchmark hoặc cluster thật.")

    doc.add_heading("2. Kiến trúc và luồng hoạt động", level=1)
    add_callout(
        doc,
        "Luồng tổng quát",
        "CSV → Validation → HDFS Raw → Bronze Delta → Silver Delta → Quality Gate → Gold → Power BI/MongoDB",
        fill="EAF2F8",
    )
    doc.add_heading("2.1 Pipeline kiểm tra nhanh", level=2)
    for text in [
        "Người dùng chạy `project_cli.py check`.",
        "CLI gọi `validate_input.py` để kiểm tra schema và business rules.",
        "CLI gọi `generate_portfolio_report.py` để tái tạo business insights.",
        "Cuối cùng CLI chạy pytest và trả nguyên exit code cho CI.",
    ]:
        add_number(doc, text)
    doc.add_heading("2.2 Pipeline Lakehouse", level=2)
    add_table(
        doc,
        ["Bước", "File/hàm chính", "Input", "Output"],
        [
            ("Ingestion", "SparkEcommerceAnalysis.main", "CSV trên HDFS", "Raw DataFrame"),
            ("Bronze", "save_delta", "Raw DataFrame", "Bronze Delta"),
            ("Silver", "chuẩn kiểu + business rules", "Bronze/raw", "Silver Delta"),
            (
                "Quality",
                "validate_raw_schema / validate_silver_data",
                "Raw/Silver",
                "Pass hoặc exception",
            ),
            ("Star schema", "build_dim_* / build_fact_sales", "Silver", "Fact + dimensions"),
            ("Data marts", "groupBy + aggregation", "Silver", "10 bảng tổng hợp"),
            ("Serving", "Hive/Thrift và InsertMongoDB", "Gold Delta", "Power BI/MongoDB"),
        ],
        [1050, 2550, 1800, 3960],
    )
    doc.add_heading("2.3 Mô hình dữ liệu Gold", level=2)
    doc.add_paragraph(
        "Grain của fact_sales là một sản phẩm trong một đơn hàng. DateKey dùng yyyyMMdd; các dimension còn lại dùng surrogate key ổn định."
    )
    add_bullet(doc, "dim_date: ngày, tháng, quý, năm.")
    add_bullet(doc, "dim_customer: khách hàng, giới tính, phân khúc.")
    add_bullet(doc, "dim_product: category, sub-category, sản phẩm.")
    add_bullet(doc, "dim_location, dim_shipping, dim_payment, dim_order_status.")

    doc.add_heading("3. Đánh giá kiến trúc hiện tại", level=1)
    add_table(
        doc,
        ["Mức độ", "Vấn đề", "Ảnh hưởng", "Cách cải thiện"],
        [
            (
                "High",
                "SparkEcommerceAnalysis.py khoảng 882 dòng",
                "Khó test, đọc và thay đổi",
                "Tách ingestion, transforms, quality, gold builders và orchestration",
            ),
            (
                "High",
                "Chưa chạy integration test vì thiếu HDFS/PySpark/Delta/PyMongo",
                "Chưa chứng minh end-to-end trên môi trường hiện tại",
                "Dựng môi trường đầy đủ hoặc Docker Compose rồi smoke test",
            ),
            (
                "Medium",
                "Một số logging vẫn dùng print",
                "Khó lọc log và quan sát production",
                "Chuẩn hóa logging có stage/run_id",
            ),
            (
                "Medium",
                "Full overwrite",
                "Không tối ưu khi dữ liệu tăng",
                "Thêm incremental MERGE, watermark và partition",
            ),
            (
                "Medium",
                "Derby metastore local",
                "Không chạy nhiều tiến trình",
                "Dùng external Hive Metastore",
            ),
            (
                "Low",
                "Hai file PBIX tên chưa chuẩn",
                "GitHub thiếu tính chuyên nghiệp",
                "Chỉ giữ dashboard cuối và ảnh preview",
            ),
            (
                "Low",
                "Ruff đã sạch tại lần kiểm tra gần nhất",
                "Cần duy trì sau mỗi thay đổi",
                "Giữ lint trong CI, không che lỗi bằng noqa",
            ),
        ],
        [1050, 2650, 2600, 3060],
    )
    add_callout(
        doc,
        "Nhận định",
        "Kiến trúc phù hợp portfolio local và thể hiện tốt năng lực Data Engineering. Điểm cần ưu tiên tiếp theo là modular hóa pipeline Spark và chứng minh integration run.",
    )

    doc.add_heading("4. Technical debt và code smell", level=1)
    add_table(
        doc,
        ["Technical debt", "Impact", "Effort", "Risk", "Priority"],
        [
            ("Pipeline Spark nguyên khối", "Cao", "Cao", "Trung bình", "P1"),
            ("Chưa incremental/partition", "Cao khi dữ liệu lớn", "Cao", "Trung bình", "P2"),
            ("Thiếu integration environment", "Cao", "Trung bình", "Thấp", "P1"),
            ("Print thay logging hoàn chỉnh", "Trung bình", "Thấp", "Thấp", "P1"),
            ("SQL Server legacy hard-coded path", "Trung bình", "Thấp", "Thấp", "P2"),
            ("Artifact runtime trong SourceCode", "Thấp", "Thấp", "Thấp", "P2"),
        ],
        [3400, 1500, 1300, 1400, 1760],
    )
    doc.add_heading("Nguyên tắc xử lý", level=2)
    add_bullet(doc, "Ưu tiên Impact cao + Effort thấp trước.")
    add_bullet(doc, "Không refactor lớn nếu không có test bảo vệ hành vi.")
    add_bullet(doc, "Không dùng noqa/type ignore để che nguyên nhân thật.")
    add_bullet(doc, "Không xóa dữ liệu, PBIX hoặc SQL legacy nếu chưa xác nhận giá trị sử dụng.")

    doc.add_heading("5. Kế hoạch cải tiến theo giai đoạn", level=1)
    phases = [
        (
            "Phase 1 — Safe Cleanup",
            "Sửa Ruff, tên file, import, dead code, docstring/comment tiếng Việt và logging cơ bản.",
        ),
        (
            "Phase 2 — Structural Refactoring",
            "Tách Spark pipeline thành modules: ingestion.py, transforms.py, dimensions.py, marts.py, quality.py, pipeline.py.",
        ),
        (
            "Phase 3 — Reliability",
            "Thêm Spark local tests, integration smoke test, retry có giới hạn, run_id và báo cáo quality metrics.",
        ),
        (
            "Phase 4 — Performance",
            "Partition Gold, incremental MERGE, giảm action count, persist đúng điểm và MongoDB Spark Connector khi scale.",
        ),
        (
            "Phase 5 — Documentation",
            "Ảnh dashboard, demo GIF/video, benchmark có điều kiện, runbook và kiến trúc production target.",
        ),
    ]
    for title, detail in phases:
        doc.add_heading(title, level=2)
        doc.add_paragraph(detail)

    doc.add_heading("6. Những thay đổi đã thực hiện", level=1)
    add_table(
        doc,
        ["Nhóm", "Thay đổi", "Lợi ích"],
        [
            (
                "Configuration",
                "config.py + biến môi trường + .env.example",
                "Không sửa source khi chuyển máy",
            ),
            ("Quality", "data_quality.py và Spark quality gate", "Dừng sớm khi dữ liệu sai"),
            ("Serving", "MongoDB đọc trực tiếp Delta theo batch", "Một nguồn sự thật, giảm RAM"),
            (
                "Automation",
                "project_cli.py với doctor/check/report/pipeline/mongodb/thrift",
                "Một entry point dễ dùng",
            ),
            ("Analytics", "portfolio_metrics.py + report generator", "KPI tái tạo được từ dữ liệu"),
            (
                "Testing",
                "12 test hiện đạt",
                "Bảo vệ quality, KPI, CLI, cấu hình, môi trường và analytics",
            ),
            ("CI", "GitHub Actions + Ruff + pytest + data contract", "Phát hiện regression sớm"),
            (
                "Docs",
                "README và docs kiến trúc/data dictionary/CV/insights",
                "Dễ đánh giá trên GitHub",
            ),
        ],
        [1600, 4960, 2800],
    )

    doc.add_heading("7. Data quality, testing và CI", level=1)
    doc.add_heading("7.1 Quality rules", level=2)
    for rule in [
        "Có đủ các cột bắt buộc.",
        "Quantity > 0; Unit_Price ≥ 0.",
        "Discount nằm trong [0, 1].",
        "Shipping_Days ≥ 0.",
        "Silver không được rỗng sau làm sạch.",
        "Order_ID được phép lặp vì một đơn có thể có nhiều dòng sản phẩm.",
    ]:
        add_bullet(doc, rule)
    doc.add_heading("7.2 Kết quả kiểm tra hiện tại", level=2)
    add_table(
        doc,
        ["Hạng mục", "Kết quả", "Ghi chú"],
        [
            ("Unit tests", "12/12 PASS", "Chạy bằng pytest trên runtime được cung cấp"),
            ("Dataset", "10.000 dòng / 26 cột", "Quality gate đầu vào đạt ở lần kiểm tra trước"),
            (
                "Ruff",
                "PASS",
                "Không còn import dư, undefined name hoặc lỗi style được chọn",
            ),
            ("Integration", "Chưa chạy", "Môi trường hiện tại thiếu HDFS, PySpark, Delta, PyMongo"),
        ],
        [2200, 2200, 4960],
    )
    add_callout(
        doc,
        "Điều kiện hoàn tất",
        "Không coi dự án end-to-end đã được xác minh cho đến khi Spark, HDFS, Delta, Hive/Thrift và MongoDB đều chạy qua smoke test.",
        fill="FFF0F0",
        color=RED,
    )

    doc.add_heading("8. Tối ưu hiệu năng và bảo mật", level=1)
    doc.add_heading("8.1 Tối ưu đã thực hiện", level=2)
    add_table(
        doc,
        ["Trước", "Sau", "Tác động"],
        [
            (
                "MongoDB tạo list toàn bộ documents",
                "toLocalIterator + batch",
                "Giảm peak memory từ O(n) xuống gần O(batch_size)",
            ),
            (
                "CSV riêng cho MongoDB",
                "Đọc trực tiếp Silver/Gold Delta",
                "Loại nguồn dữ liệu trùng và nguy cơ lệch số",
            ),
            ("Pipeline dừng chờ input", "Mặc định tự kết thúc", "Phù hợp CI/orchestrator"),
            (
                "KPI có thể tính sai trung bình tỷ lệ",
                "Tính từ tổng tử số/mẫu số",
                "KPI đúng theo business grain",
            ),
        ],
        [3100, 3100, 3160],
    )
    doc.add_heading("8.2 Tối ưu đề xuất", level=2)
    add_bullet(
        doc,
        "Giảm Spark actions lặp: gom quality metrics vào một aggregation thay vì nhiều count riêng.",
    )
    add_bullet(doc, "Partition fact/marts theo Year/Month khi dữ liệu đủ lớn.")
    add_bullet(
        doc, "Dùng Delta MERGE cho incremental load; chỉ overwrite dimension nhỏ khi phù hợp."
    )
    add_bullet(doc, "Benchmark trước/sau trên cùng dataset, cấu hình Spark và phần cứng.")
    doc.add_heading("8.3 Security", level=2)
    add_bullet(doc, "URI và cấu hình được lấy từ biến môi trường; .env bị loại khỏi Git.")
    add_bullet(doc, "Thrift command dùng list argument, không ghép shell string.")
    add_bullet(doc, "Không log secret; production nên dùng secret manager và authentication/TLS.")
    add_bullet(
        doc, "Dataset mẫu không nên chứa PII thực tế; cần masking và access control khi production."
    )

    doc.add_heading("9. Cách cài đặt và chạy dự án", level=1)
    doc.add_heading("9.1 Chuẩn bị", level=2)
    for step in [
        "Cài Python 3.10/3.11, Java 11/17, Hadoop/HDFS và MongoDB nếu sử dụng đầy đủ.",
        "Tạo virtual environment và cài `requirements.txt`.",
        "Thiết lập biến môi trường theo `.env.example`.",
        "Chạy `python SourceCode/project_cli.py doctor` để kiểm tra dependency.",
    ]:
        add_number(doc, step)
    doc.add_heading("9.2 Các lệnh chính", level=2)
    add_table(
        doc,
        ["Lệnh", "Mục đích"],
        [
            ("python SourceCode/project_cli.py doctor", "Kiểm tra Java, HDFS và Python packages"),
            ("python SourceCode/project_cli.py check", "Validate data, sinh report, chạy tests"),
            ("python SourceCode/project_cli.py pipeline", "Chạy Bronze–Silver–Gold"),
            ("python SourceCode/project_cli.py mongodb", "Đồng bộ Delta sang MongoDB"),
            ("python SourceCode/project_cli.py thrift", "Mở endpoint cho Power BI"),
        ],
        [5000, 4360],
    )
    doc.add_heading("9.3 Trình tự demo CV", level=2)
    add_number(doc, "Mở README và giới thiệu bài toán, kiến trúc trong 30–45 giây.")
    add_number(doc, "Chạy doctor/check để chứng minh khả năng tái lập và kiểm thử.")
    add_number(doc, "Trình bày sơ đồ Medallion + star schema và giải thích grain.")
    add_number(doc, "Mở Power BI dashboard, nêu 3 insights và một khuyến nghị kinh doanh.")
    add_number(doc, "Kết thúc bằng hạn chế hiện tại và roadmap production.")

    doc.add_heading("10. Kết quả, hạn chế và bước tiếp theo", level=1)
    doc.add_heading("10.1 Kết quả có thể kiểm chứng", level=2)
    add_bullet(doc, "12 test hiện tại đều pass.")
    add_bullet(doc, "Báo cáo business insights tái tạo từ dataset, không nhập KPI thủ công.")
    add_bullet(doc, "CLI, CI, data contract và cấu hình tập trung đã sẵn sàng.")
    add_bullet(doc, "Kiến trúc và data dictionary đã được tài liệu hóa.")
    doc.add_heading("10.2 Hạn chế", level=2)
    add_bullet(doc, "Chưa có integration run trong môi trường hiện tại.")
    add_bullet(doc, "Spark pipeline còn lớn và cần refactor có kiểm soát.")
    add_bullet(doc, "Chưa có incremental load, orchestration, alerting và benchmark.")
    add_bullet(doc, "PBIX cần chuẩn hóa tên và bổ sung ảnh preview trong README.")
    doc.add_heading("10.3 Thứ tự ưu tiên", level=2)
    add_number(doc, "Duy trì Ruff sạch và chạy toàn bộ 12 tests sau mỗi thay đổi.")
    add_number(doc, "Thiết lập HDFS/PySpark/Delta/PyMongo rồi chạy smoke test end-to-end.")
    add_number(doc, "Tách pipeline Spark theo layer, mỗi thay đổi kèm regression test.")
    add_number(doc, "Thêm ảnh Power BI và chuẩn hóa repository public.")
    add_number(doc, "Sau đó mới triển khai incremental load và orchestration.")

    doc.add_heading("11. Cách trình bày trong CV và phỏng vấn", level=1)
    doc.add_heading("Tên dự án", level=2)
    add_callout(
        doc,
        "Khuyến nghị",
        "GlobalCart Intelligence — End-to-End Ecommerce Lakehouse & BI Platform",
        fill="EAF2F8",
    )
    doc.add_heading("Bullet CV tiếng Việt", level=2)
    add_bullet(
        doc,
        "Thiết kế Medallion Lakehouse xử lý 10.000 giao dịch bằng PySpark, HDFS và Delta Lake, chuẩn hóa dữ liệu từ Raw/Bronze đến Silver/Gold.",
    )
    add_bullet(
        doc,
        "Xây dựng Kimball star schema gồm 1 fact, 7 dimensions và 10 data marts cho Power BI, phục vụ phân tích doanh thu, lợi nhuận, khách hàng, sản phẩm và fulfillment.",
    )
    add_bullet(
        doc,
        "Triển khai data quality gate, cấu hình theo môi trường, batch MongoDB synchronization, unit tests và CI checks từ một Delta source of truth.",
    )
    doc.add_heading("Business insights để kể", level=2)
    add_bullet(doc, "Doanh thu khoảng 5,28 triệu USD; biên lợi nhuận có trọng số 27,21%.")
    add_bullet(doc, "Electronics đóng góp gần 64% doanh thu, thể hiện rủi ro tập trung ngành hàng.")
    add_bullet(
        doc,
        "Returned + Cancelled khoảng 27,65%, là cơ hội drill-down theo sản phẩm, quốc gia và vận chuyển.",
    )
    add_callout(
        doc,
        "Lưu ý",
        "Chỉ nêu số liệu hiệu năng nếu đã benchmark. Trung thực về môi trường local tạo ấn tượng tốt hơn việc gọi dự án là production-ready khi chưa có bằng chứng.",
        fill="FFF8E8",
        color=GOLD,
    )

    doc.add_heading("12. Checklist bàn giao", level=1)
    checks = [
        "[ ] Duy trì trạng thái Ruff All checks passed.",
        "[ ] Cài HDFS, PySpark, Delta Lake, PyMongo theo requirements.",
        "[ ] Chạy doctor, check và pipeline từ đầu đến cuối.",
        "[ ] Xác minh số dòng Bronze/Silver/Fact và tất cả foreign keys.",
        "[ ] Xác minh MongoDB collection count và index.",
        "[ ] Mở Power BI qua Thrift, refresh thành công.",
        "[ ] Đổi tên PBIX cuối thành GlobalCart_Intelligence_Dashboard.pbix.",
        "[ ] Thêm 2–3 ảnh dashboard vào docs/images và README.",
        "[ ] Đổi tên repository thành globalcart-intelligence.",
        "[ ] Kiểm tra repository không chứa secret, cache hoặc runtime database.",
        "[ ] Gắn GitHub URL vào CV và luyện phần trình bày 3 phút.",
    ]
    for item in checks:
        add_bullet(doc, item)

    doc.add_heading("Phụ lục — Danh sách file trọng tâm", level=1)
    add_table(
        doc,
        ["File", "Vai trò"],
        [
            ("README.md", "Hướng dẫn chính và landing page GitHub"),
            ("SourceCode/project_cli.py", "Entry point thống nhất"),
            ("SourceCode/SparkEcommerceAnalysis.py", "Pipeline Lakehouse chính"),
            ("SourceCode/data_quality.py", "Data contract và quality rules"),
            ("SourceCode/InsertMongoDB.py", "Serving sang MongoDB"),
            ("SourceCode/portfolio_metrics.py", "KPI có thể kiểm thử"),
            (".github/workflows/quality.yml", "CI lint/test/data validation"),
            ("docs/BUSINESS_INSIGHTS.md", "Báo cáo insight tái tạo từ dữ liệu"),
            ("tests/", "Regression tests cho logic quan trọng"),
        ],
        [3900, 5460],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
